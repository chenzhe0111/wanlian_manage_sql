#!/usr/bin/env python3
"""用本地《1014目标管理数据.xlsx》回写飞书目标表（跳过 QuickBI）。

sheet 映射：
  调度1          → 1074dd 月实际（46b）
  月度新老货主   → 1074dd 第七部分（49）
  司机           → 41UJIa
  周度汇总       → 66mgnV（近两周）
  03月度周度     → QbfdXI（03d：含上月同期 / 较上月同期）

as_of = 调度1 10 月「统计天数」对应日（13 → 2026-10-13）。
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SHEET_URL = "https://wanlianyida.feishu.cn/sheets/L91EsEtAohwjhStD4y6cYLCGnVe"
XLSX = Path("/Users/chenzhe/Desktop/1014目标管理数据.xlsx")

from refresh_sheet_actuals import build_writes as build_actuals_writes  # noqa: E402
from refresh_sheet_actuals import write_sheet as write_actuals  # noqa: E402
from refresh_sheet_driver import build_writes as build_driver_writes  # noqa: E402
from refresh_sheet_driver import write_sheet as write_driver  # noqa: E402
from refresh_sheet_funnel import INIT_ORDER, METRIC_ORDER, SHIPPER_ORDER  # noqa: E402
from refresh_sheet_newold import build_writes as build_newold_writes  # noqa: E402
from refresh_sheet_newold import write_sheet as write_newold  # noqa: E402
from refresh_sheet_weekly import build_table as build_weekly_table  # noqa: E402
from refresh_sheet_weekly import write_sheet as write_weekly  # noqa: E402


def _lark(*args: str) -> dict:
    cmd = ["lark-cli", "sheets", *args, "--url", SHEET_URL, "--as", "user", "--format", "json"]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    raw = proc.stdout or proc.stderr
    i = raw.find("{")
    if proc.returncode != 0 or i < 0:
        raise SystemExit(raw[-2000:] or f"lark-cli failed {proc.returncode}")
    data = json.loads(raw[i:])
    if not data.get("ok", True):
        raise SystemExit(json.dumps(data.get("error") or data, ensure_ascii=False)[:2000])
    return data


def _to_date(v) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v)[:10]
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


def _num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_actuals(ws, month: date) -> tuple[date, list[dict]]:
    rows: list[dict] = []
    day_n = None
    mon_key = month.strftime("%Y%m")
    for r in range(2, ws.max_row + 1):
        mon = _to_date(ws.cell(r, 1).value)
        if mon is None or mon.year != month.year or mon.month != month.month:
            continue
        day_n = int(ws.cell(r, 4).value)
        rows.append(
            {
                "month": mon_key,
                "key1": str(ws.cell(r, 2).value or "").strip(),
                "key2": str(ws.cell(r, 3).value or "").strip(),
                "month_num": ws.cell(r, 5).value,
                "avg_num": ws.cell(r, 6).value,
                "max_num": ws.cell(r, 7).value,
                "gtv_1": ws.cell(r, 9).value,
                "avg_cost": ws.cell(r, 11).value,
            }
        )
    if not rows or not day_n:
        raise SystemExit(f"调度1 缺少 {month.strftime('%Y-%m')} 行")
    return date(month.year, month.month, day_n), rows


def load_newold_rows(ws, month: date) -> list[dict]:
    mon_dash = month.strftime("%Y-%m")
    out: list[dict] = []
    for r in range(2, ws.max_row + 1):
        mon = _to_date(ws.cell(r, 1).value)
        if mon is None or mon.year != month.year or mon.month != month.month:
            continue
        dim = str(ws.cell(r, 3).value or "").strip()
        out.append(
            {
                "month": mon_dash,
                "index2": dim,
                "fahuo_num": ws.cell(r, 5).value,
                "chengjiao_num": ws.cell(r, 6).value,
                "chengjiao_yundan": ws.cell(r, 8).value,  # 万单
                "avg_num": ws.cell(r, 10).value,
                "zhuce_num": ws.cell(r, 11).value,
                "renzheng_num": ws.cell(r, 12).value,
            }
        )
    if not out:
        raise SystemExit(f"月度新老货主缺少 {mon_dash}")
    return out


def load_driver_row(ws, month: date) -> tuple[str, dict]:
    for r in range(2, ws.max_row + 1):
        mon = _to_date(ws.cell(r, 2).value)
        if mon is None or mon.year != month.year or mon.month != month.month:
            continue
        pt = _to_date(ws.cell(r, 1).value)
        as_of = pt.isoformat() if pt else month.isoformat()
        row = {
            "tms_driver": ws.cell(r, 6).value,
            "tms_avg": ws.cell(r, 7).value,
            "tms_num": ws.cell(r, 10).value,
            "cuohe_driver": ws.cell(r, 12).value,
            "cuohe_avg": ws.cell(r, 13).value,
            "cuohe_num": ws.cell(r, 16).value,
            "wanghuo_driver": ws.cell(r, 18).value,
            "wanghuo_avg": ws.cell(r, 19).value,
            "wanghuo_num": ws.cell(r, 22).value,
        }
        return as_of, row
    raise SystemExit(f"司机缺少 {month.strftime('%Y-%m')}")


def load_weekly_blocks(ws, as_of: date):
    last_wed = as_of - timedelta(days=(as_of.weekday() - 2) % 7)
    targets = [last_wed, last_wed - timedelta(days=7)]
    by_week: dict[date, list[dict]] = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        ws_d = _to_date(ws.cell(r, 1).value)
        if ws_d is None:
            continue
        by_week[ws_d].append(
            {
                "data_index": str(ws.cell(r, 2).value or "").strip(),
                "index1": str(ws.cell(r, 3).value or "").strip(),
                "index2": str(ws.cell(r, 4).value or "").strip(),
                "num": ws.cell(r, 5).value,
            }
        )
    weeks = []
    missing = []
    for ws_d in targets:
        if ws_d in by_week:
            weeks.append((ws_d, ws_d - timedelta(days=1), by_week[ws_d]))
        else:
            missing.append(ws_d)
    return weeks, missing, targets


def load_funnel_blocks_03d(ws, as_of: date):
    month_start = as_of.replace(day=1)
    last_wed = as_of - timedelta(days=(as_of.weekday() - 2) % 7)
    week_targets = [last_wed, last_wed - timedelta(days=7)]

    by_period: dict[tuple[str, date], list[dict]] = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        grain = str(ws.cell(r, 2).value or "").strip()
        ps = _to_date(ws.cell(r, 3).value)
        if not grain or ps is None:
            continue
        by_period[(grain, ps)].append(
            {
                "key": str(ws.cell(r, 1).value or "").strip(),
                "type": str(ws.cell(r, 4).value or "").strip(),
                "new_old": str(ws.cell(r, 5).value or "").strip(),
                "index_type": str(ws.cell(r, 6).value or "").strip(),
                "num": ws.cell(r, 7).value,
                "prev_num": ws.cell(r, 8).value,
                "yoy": ws.cell(r, 9).value,
            }
        )

    blocks = []
    missing = []
    if ("月度", month_start) in by_period:
        blocks.append(("月度", month_start, by_period[("月度", month_start)]))
    else:
        missing.append(f"月度{month_start.isoformat()}")
    for ws_d in week_targets:
        if ("周度", ws_d) in by_period:
            blocks.append(("周度", ws_d, by_period[("周度", ws_d)]))
        else:
            missing.append(f"周度{ws_d.isoformat()}")
    return blocks, missing, week_targets


def _sort_funnel(r: dict) -> tuple:
    return (
        INIT_ORDER.get(str(r.get("type") or "").strip(), 99),
        SHIPPER_ORDER.get(str(r.get("new_old") or "").strip(), 99),
        METRIC_ORDER.get(str(r.get("index_type") or "").strip(), 99),
        str(r.get("key") or ""),
    )


def build_funnel_grid_03d(blocks) -> list[list[dict]]:
    header = [
        "月周标识",
        "主键",
        "周期起始日",
        "举措",
        "新老货主",
        "指标类型",
        "指标值",
        "上月同期指标值",
        "较上月同期",
    ]
    grid: list[list[dict]] = [[{"value": h} for h in header]]

    def cell(v):
        # lark cells-set 不接受 JSON null；空值用 ""
        if v is None:
            return {"value": ""}
        return {"value": v}

    def fmt(v):
        n = _num(v)
        if n is None:
            return ""
        if isinstance(n, float) and n == int(n):
            return int(n)
        return round(n, 6) if isinstance(n, float) else n

    for grain, period_start, rows in blocks:
        ps = period_start.isoformat()
        for r in sorted(rows, key=_sort_funnel):
            grid.append(
                [
                    cell(grain),
                    cell(str(r.get("key") or "").strip()),
                    cell(ps),
                    cell(str(r.get("type") or "").strip()),
                    cell(str(r.get("new_old") or "").strip()),
                    cell(str(r.get("index_type") or "").strip()),
                    cell(fmt(r.get("num"))),
                    cell(fmt(r.get("prev_num"))),
                    cell(fmt(r.get("yoy"))),
                ]
            )
    return grid


def write_funnel_03d(grid: list[list[dict]], title: str) -> None:
    """同 write_funnel，但列扩到 I（03d 同期列）。"""
    n = len(grid)
    ops = [
        {
            "shortcut": "+cells-clear",
            "input": {"sheet_id": "QbfdXI", "range": "A1:I2000", "scope": "all"},
        },
        {
            "shortcut": "+cells-set",
            "input": {
                "sheet_id": "QbfdXI",
                "range": "A1",
                "cells": [
                    [
                        {
                            "value": title,
                            "cell_styles": {
                                "font_weight": "bold",
                                "font_size": 13,
                                "background_color": "#FFF2CC",
                            },
                        }
                    ]
                ],
            },
        },
        {
            "shortcut": "+cells-merge",
            "input": {"sheet_id": "QbfdXI", "range": "A1:I1", "merge_type": "all"},
        },
        {
            "shortcut": "+cells-set",
            "input": {
                "sheet_id": "QbfdXI",
                "range": f"A2:I{n + 1}",
                "cells": [
                    [
                        {
                            **c,
                            "cell_styles": {
                                **(
                                    {
                                        "background_color": "#BDD7EE",
                                        "font_weight": "bold",
                                        "horizontal_alignment": "center",
                                    }
                                    if i == 0
                                    else {}
                                ),
                                "vertical_alignment": "middle",
                                "font_size": 11,
                            },
                        }
                        for c in row
                    ]
                    for i, row in enumerate(grid)
                ],
            },
        },
    ]
    payload = ROOT / "logs" / "sheet_funnel_ops.json"
    payload.parent.mkdir(parents=True, exist_ok=True)
    payload.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
    rel = payload.relative_to(ROOT).as_posix()
    # reuse lark path from funnel module via local call
    cmd = [
        "lark-cli",
        "sheets",
        "+batch-update",
        "--yes",
        "--operations",
        f"@./{rel}",
        "--url",
        SHEET_URL,
        "--as",
        "user",
        "--format",
        "json",
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, cwd=ROOT)
    raw = proc.stdout or proc.stderr
    i = raw.find("{")
    if proc.returncode != 0 or i < 0:
        raise SystemExit(raw[-2000:] or "funnel write failed")
    data = json.loads(raw[i:])
    if not data.get("ok", True):
        raise SystemExit(json.dumps(data.get("error") or data, ensure_ascii=False)[:2000])


def main() -> int:
    if not XLSX.exists():
        raise SystemExit(f"找不到 {XLSX}")

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    target_month = date(2026, 10, 1)
    as_of, overall = load_actuals(wb["调度1"], target_month)
    as_of_s = as_of.isoformat()
    month_key = "202610"
    month_dash = "2026-10"
    print(f"[{datetime.now():%F %T}] excel={XLSX.name} as_of={as_of_s}")

    os.chdir(ROOT)

    # 1) 整合表月实际
    actuals = build_actuals_writes(overall, month_key)
    print(f"actuals cells={len(actuals)} N5={actuals.get('N5')} N65={actuals.get('N65')} N73={actuals.get('N73')}")
    write_actuals(actuals, as_of_s)
    print("✓ 1074dd 月实际已写")

    # 2) 第七部分新老（49）
    newold_rows = load_newold_rows(wb["月度新老货主"], target_month)
    newold = build_newold_writes(newold_rows, month_dash)
    print(
        f"newold cells={len(newold)} N103={newold.get('N103')} N120={newold.get('N120')} "
        f"N157={newold.get('N157')} N158={newold.get('N158')}"
    )
    write_newold(newold, as_of_s)
    print("✓ 1074dd 第七部分新老已写")

    # 3) 司机
    drv_asof, drv_row = load_driver_row(wb["司机"], target_month)
    driver = build_driver_writes(drv_row, month_dash)
    print("driver writes", json.dumps(driver, ensure_ascii=False))
    write_driver(driver, as_of=drv_asof)
    print("✓ 41UJIa 司机已写")

    # 4) 周度详细
    weeks, miss_w, week_targets = load_weekly_blocks(wb["周度汇总"], as_of)
    if weeks:
        got_w = "、".join(w.strftime("%m/%d") for w, _, _ in weeks)
        wtitle = (
            f"周度详细数据｜最近两周目标 {week_targets[0].strftime('%m/%d')}&"
            f"{week_targets[1].strftime('%m/%d')}｜已写入 {got_w}｜刷新日 {as_of_s}｜本地Excel1014"
        )
        if miss_w:
            wtitle += "｜暂无 " + "、".join(m.strftime("%m/%d") for m in miss_w)
        write_weekly(build_weekly_table(weeks), wtitle)
        print(f"✓ 66mgnV 周度已写 weeks={[w[0].isoformat() for w in weeks]} miss={miss_w}")
    else:
        print(f"⚠ 周度跳过 miss={miss_w}")

    # 5) 漏斗 03d
    blocks, miss_f, week_targets = load_funnel_blocks_03d(wb["03月度周度"], as_of)
    if not blocks:
        raise SystemExit(f"漏斗无可用块: missing={miss_f}")
    got = "、".join(("月" if g == "月度" else "周") + p.strftime("%m/%d") for g, p, _ in blocks)
    title = (
        f"月度周度漏斗(03d)｜本月 {as_of.strftime('%Y-%m')}｜近两周 "
        f"{week_targets[0].strftime('%m/%d')}&{week_targets[1].strftime('%m/%d')}"
        f"｜已写入 {got}｜刷新日 {as_of_s}｜本地Excel1014"
    )
    if miss_f:
        title += "｜暂无 " + "、".join(miss_f)
    write_funnel_03d(build_funnel_grid_03d(blocks), title)
    print(
        f"✓ QbfdXI 漏斗(03d)已写 blocks="
        f"{[(g, p.isoformat(), len(r)) for g, p, r in blocks]} miss={miss_f}"
    )

    # 回读核对
    print("\n--- verify ---")
    v1 = _lark("+cells-get", "--sheet-id", "1074dd", "--range", "O1,N5,N65,N73,N103,N120,N157")
    print("1074dd:", json.dumps(v1.get("data") or v1, ensure_ascii=False)[:900])
    v2 = _lark("+cells-get", "--sheet-id", "41UJIa", "--range", "L1,K7,K8,K10")
    print("41UJIa:", json.dumps(v2.get("data") or v2, ensure_ascii=False)[:500])
    v3 = _lark("+cells-get", "--sheet-id", "QbfdXI", "--range", "A1:I3")
    print("QbfdXI:", json.dumps(v3.get("data") or v3, ensure_ascii=False)[:700])
    v4 = _lark("+cells-get", "--sheet-id", "66mgnV", "--range", "A1:E3")
    print("66mgnV:", json.dumps(v4.get("data") or v4, ensure_ascii=False)[:500])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
